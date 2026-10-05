import os
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DOCS_DIR = PROJECT_ROOT / "docs"
SCREENSHOTS_DIR = DOCS_DIR / "screenshots"
OUTPUT_PATH = DOCS_DIR / "KBM_TenderScout_User_Guide.docx"

# Color Palette
COLOR_NAVY = RGBColor(15, 38, 66)       # #0F2642
COLOR_BLUE = RGBColor(30, 64, 175)      # #1E40AF
COLOR_ACCENT = RGBColor(37, 99, 235)    # #2563EB
COLOR_TEXT = RGBColor(30, 41, 59)       # #1E293B
COLOR_MUTED = RGBColor(100, 116, 139)   # #64748B
COLOR_GREEN = RGBColor(22, 101, 52)     # #166534

HEX_NAVY = "0F2642"
HEX_LIGHT_BLUE = "EFF6FF"
HEX_LIGHT_GRAY = "F8FAFC"
HEX_BORDER = "CBD5E1"
HEX_ACCENT = "2563EB"
HEX_WHITE = "FFFFFF"

def set_cell_shading(cell, color_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_callout(doc, text_paragraphs, title="NOTE / ملاحظة هامة"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.4)
    set_cell_shading(cell, HEX_LIGHT_BLUE)
    set_cell_margins(cell, top=160, bottom=160, left=220, right=200)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{HEX_ACCENT}"/>
            <w:top w:val="none"/>
            <w:right w:val="none"/>
            <w:bottom w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_title = p.add_run(f"📌 {title}\n")
    run_title.font.bold = True
    run_title.font.name = "Segoe UI"
    run_title.font.size = Pt(10.5)
    run_title.font.color.rgb = COLOR_BLUE
    
    for i, line in enumerate(text_paragraphs):
        if i > 0:
            p = cell.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
        run = p.add_run(line)
        run.font.name = "Segoe UI"
        run.font.size = Pt(9.5)
        run.font.color.rgb = COLOR_TEXT

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(4)
    p_after.paragraph_format.space_after = Pt(6)

def style_heading_1(p):
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    for run in p.runs:
        run.font.name = "Segoe UI"
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = COLOR_NAVY

def style_heading_2(p):
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    for run in p.runs:
        run.font.name = "Segoe UI"
        run.font.size = Pt(12.5)
        run.font.bold = True
        run.font.color.rgb = COLOR_BLUE

def style_body(p, space_after=4):
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    for run in p.runs:
        run.font.name = "Segoe UI"
        run.font.size = Pt(10)
        run.font.color.rgb = COLOR_TEXT

def add_image_with_caption(doc, image_name, caption_text, width=Inches(6.2)):
    img_path = SCREENSHOTS_DIR / image_name
    if img_path.exists():
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        run = p_img.add_run()
        run.add_picture(str(img_path), width=width)
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(0)
        p_cap.paragraph_format.space_after = Pt(10)
        run_cap = p_cap.add_run(caption_text)
        run_cap.font.name = "Segoe UI"
        run_cap.font.size = Pt(9)
        run_cap.font.italic = True
        run_cap.font.color.rgb = COLOR_MUTED
    else:
        print(f"Warning: Image not found {img_path}")

def build_guide():
    doc = Document()
    
    # Page setup - Margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
        # Header & Footer
        header = section.header
        p_head = header.paragraphs[0]
        p_head.text = "KBM TenderScout™ | User Guide & Operational Manual"
        p_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_head.runs[0].font.name = "Segoe UI"
        p_head.runs[0].font.size = Pt(8.5)
        p_head.runs[0].font.color.rgb = COLOR_MUTED
        
        footer = section.footer
        p_foot = footer.paragraphs[0]
        p_foot.text = "Confidential - KBM Internal Operations | Live Platform: https://kbm-tenderscout.azurewebsites.net"
        p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_foot.runs[0].font.name = "Segoe UI"
        p_foot.runs[0].font.size = Pt(8)
        p_foot.runs[0].font.color.rgb = COLOR_MUTED

    # ================= COVER / TITLE BLOCK =================
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(2)
    run_org = p_title.add_run("KHALID BOSHNAK & MOHAMMAD (KBM) TECHNOLOGY PIPELINE\n")
    run_org.font.name = "Segoe UI"
    run_org.font.size = Pt(11)
    run_org.font.bold = True
    run_org.font.color.rgb = COLOR_ACCENT

    run_main_title = p_title.add_run("KBM TenderScout™ Platform\nUser Guide & Operational Manual")
    run_main_title.font.name = "Segoe UI"
    run_main_title.font.size = Pt(24)
    run_main_title.font.bold = True
    run_main_title.font.color.rgb = COLOR_NAVY

    p_ar_title = doc.add_paragraph()
    p_ar_title.paragraph_format.space_before = Pt(0)
    p_ar_title.paragraph_format.space_after = Pt(12)
    run_ar = p_ar_title.add_run("دليل المستخدم والتشغيل لمنظومة رصد وتتبع المناقصات والممارسات الحكومية والنفطية")
    run_ar.font.name = "Segoe UI"
    run_ar.font.size = Pt(13)
    run_ar.font.bold = True
    run_ar.font.color.rgb = COLOR_TEXT

    # Quick Metadata Table
    table_meta = doc.add_table(rows=5, cols=2)
    table_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Cloud Production URL:", "https://kbm-tenderscout.azurewebsites.net (Azure Hosted)"),
        ("Source Code Repository:", "https://github.com/6439144/KBM_TenderScout (GitHub Master)"),
        ("Marketing Leadership:", "Eiman Ashkanani (Marketing Manager - Government & Oil Sector)"),
        ("Automated Daily Schedule:", "08:00 AM Kuwait Time (05:00 UTC) Daily Sync"),
        ("Target Users:", "KBM Account Managers, Bidding Team, Technology Sales & Management")
    ]
    for i, (k, v) in enumerate(meta_data):
        row = table_meta.rows[i]
        c1, c2 = row.cells[0], row.cells[1]
        c1.width = Inches(2.2)
        c2.width = Inches(4.2)
        set_cell_shading(c1, HEX_LIGHT_GRAY)
        set_cell_shading(c2, HEX_LIGHT_BLUE)
        set_cell_margins(c1, top=70, bottom=70, left=120, right=120)
        set_cell_margins(c2, top=70, bottom=70, left=120, right=120)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(k)
        r1.font.bold = True
        r1.font.size = Pt(9)
        r1.font.color.rgb = COLOR_NAVY
        
        p2 = c2.paragraphs[0]
        r2 = p2.add_run(v)
        r2.font.bold = ("https" in v or "Marketing Manager" in v)
        r2.font.size = Pt(9)
        r2.font.color.rgb = COLOR_BLUE if "https" in v else COLOR_TEXT

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(8)
    p_div.paragraph_format.space_after = Pt(12)

    # ================= 1. EXECUTIVE OVERVIEW =================
    h1 = doc.add_paragraph("1. Executive Overview & System Architecture")
    style_heading_1(h1)

    p = doc.add_paragraph(
        "KBM TenderScout™ is an intelligent procurement tracking platform built specifically for KBM to monitor, "
        "analyze, and match Kuwaiti government and oil-sector public tenders and practices (ممارسات ومناقصات). "
        "The system replaces manual tracking with automated scrapers, deep natural language parsing, and strategic "
        "account manager mapping directly reflecting the updated organizational structure led by Eiman Ashkanani."
    )
    style_body(p)

    p = doc.add_paragraph(
        "The platform aggregates two authoritative state channels:\n"
        "• Central Agency for Public Tenders (CAPT / الجهاز المركزي للمناقصات العامة) - All opening and active public tenders.\n"
        "• Official Kuwait Gazette (Kuwait Al-Yawm / جريدة الكويت اليوم - ملحق المناقصات) - Full downloadable PDF supplements with extracted tender scopes, technical specifications, closing dates, pre-tender meetings, and bid bonds."
    )
    style_body(p)

    add_callout(
        doc,
        [
            "Zero Missing Data Guarantee: Incomplete entries with generic placeholder text ('LOADING...') have been completely purged.",
            "Every tender in the system contains authentic client entity names, detailed scope descriptions (e.g. Cisco core switches, IBM Maximo, enterprise servers), accurate closing dates, and official reference numbers."
        ],
        title="High Precision & Deep Scope Parsing"
    )

    # ================= 2. ACCESSING THE PLATFORM =================
    h1 = doc.add_paragraph("2. Accessing the Platform")
    style_heading_1(h1)

    p = doc.add_paragraph(
        "The application is accessible 24/7 without requiring software installation or local setups on employee machines:"
    )
    style_body(p)

    p = doc.add_paragraph(
        "1. Live Cloud Portal: https://kbm-tenderscout.azurewebsites.net\n"
        "   Hosted on Microsoft Azure App Service Free Tier (zero recurring cloud infrastructure cost).\n"
        "2. Browser Compatibility: Fully tested and responsive on Google Chrome, Microsoft Edge, Mozilla Firefox, and Apple Safari (desktop, tablet, and mobile).\n"
        "3. Local Fallback (Optional): Can be launched locally on the office network via http://127.0.0.1:8000 using the provided startup script."
    )
    style_body(p)

    add_image_with_caption(doc, "01_dashboard_overview.png", "Figure 1: KBM TenderScout Main Executive Dashboard Overview")

    # ================= 3. SYSTEM FEATURES & USER WORKFLOW =================
    h1 = doc.add_paragraph("3. Step-by-Step User Workflow & Features")
    style_heading_1(h1)

    h2 = doc.add_paragraph("3.1 KPI Metrics & Real-time Filter Bar")
    style_heading_2(h2)

    p = doc.add_paragraph(
        "The top summary section provides instantaneous situational awareness across all active opportunities:"
    )
    style_body(p)

    p = doc.add_paragraph(
        "• Total Monitored: Total count of canonical opportunities in the database (231 active tenders).\n"
        "• Strategic Matches: Tenders specifically matched to KBM core technology offerings (Cisco, IBM, HPE/Dell, Cyber, Cloud, Infrastructure).\n"
        "• Approaching Deadlines: Opportunities closing within the next 14 days, flagged in urgent amber/red.\n"
        "• Source Breakdown: Real-time counter of CAPT portal listings vs Kuwait Al-Yawm Gazette listings."
    )
    style_body(p)

    add_image_with_caption(doc, "02_kpi_and_filters.png", "Figure 2: Real-time KPI Statistics Cards and Multi-Criteria Filtering Controls")

    p = doc.add_paragraph(
        "Filtering Controls Available:\n"
        "1. Keyword Search: Type any keyword in Arabic or English (e.g., 'Cisco', 'خوادم', 'شبكات', 'PIC', 'وزارة الدفاع'). Results filter instantly.\n"
        "2. Account Manager Filter: Dropdown selector allowing each sales representative to isolate tenders relevant to their designated portfolio.\n"
        "3. Status Filter: Toggle between Active, Closing Soon, Closed, or Pending Inquiry.\n"
        "4. Source Filter: Filter by CAPT or Kuwait Al-Yawm specifically.\n"
        "5. Reset Filters: Instantly restores the default comprehensive view."
    )
    style_body(p)

    h2 = doc.add_paragraph("3.2 Strategic Opportunity Identification & Scope Badges")
    style_heading_2(h2)

    p = doc.add_paragraph(
        "The list view emphasizes high-yield opportunities by displaying technology tags and client badges directly on the card header. "
        "This allows the bidding team to scan through hundreds of tenders in seconds."
    )
    style_body(p)

    add_image_with_caption(doc, "03_strategic_tenders_list.png", "Figure 3: Strategic Tenders List with Technology Badges & Assigned Account Owners")

    h2 = doc.add_paragraph("3.3 Deep Scope Modal & Tender Requirements Analysis")
    style_heading_2(h2)

    p = doc.add_paragraph(
        "Clicking 'عرض التفاصيل والاشتراطات' (View Details) opens the comprehensive Tender Deep-Dive Modal. "
        "Unlike generic portals that only show a brief title, TenderScout provides the full operational context:"
    )
    style_body(p)

    p = doc.add_paragraph(
        "• Official Issuing Entity: (e.g. Petrochemical Industries Company PIC, Ministry of Defence, Ministry of Health).\n"
        "• Technology Scope & KBM Fit: Specific equipment models, manufacturer requirements (Cisco, IBM Maximo, Dell/HPE), and managed service expectations.\n"
        "• Commercial & Bidding Milestones: Closing date, Pre-tender technical meeting date, and Inquiry deadline.\n"
        "• Financial Requirements: Bid bond amount (الكفالة الأولية) and document purchase fee (ثمن الوثائق).\n"
        "• Submission Channel & Location: Exact tender box number, gate location, or online e-procurement link."
    )
    style_body(p)

    add_image_with_caption(doc, "04_modal_mna19_cisco.png", "Figure 4: Modal View - PIC Tender MNA/19/2026 for Cisco Managed Network Services")

    add_image_with_caption(doc, "05_modal_mna24_ibm.png", "Figure 5: Modal View - PIC Tender MNA/24/2026 for IBM Maximo Enterprise Asset Management")

    add_image_with_caption(doc, "06_modal_mod_servers.png", "Figure 6: Modal View - Ministry of Defence Practice 24181118110 for Enterprise Servers")

    h2 = doc.add_paragraph("3.4 Account Manager Portfolio Filtering & Owner Reassignment")
    style_heading_2(h2)

    p = doc.add_paragraph(
        "Every tender is automatically mapped to the appropriate KBM Account Manager according to the official portfolio allocation. "
        "Sales reps can filter by their name to view only their portfolio (e.g., Abrar Al-Qallaf for KNPC & PIC, "
        "Ahmed Habib for KOC & KOTC, Jana Al-Obaid for MOH & PIFSS, Raed Obeid for MEW & PACI, Khaled Alabdallah for MOI, "
        "and Eiman Ashkanani for Ministry of Defence and all other strategic accounts)."
    )
    style_body(p)

    add_image_with_caption(doc, "07_filtered_by_account_manager.png", "Figure 7: Filtered View for Account Manager Abrar Al-Qallaf showing Assigned Strategic Pipeline")

    p = doc.add_paragraph(
        "Reassigning Account Ownership:\n"
        "The interface includes a quick reassignment feature. If an account manager changes or an opportunity falls under another team member's domain, "
        "authorized team members can update the assigned owner directly in the database with immediate synchronization."
    )
    style_body(p)

    # ================= 4. ACCOUNT MANAGER ASSIGNMENT MATRIX =================
    h1 = doc.add_paragraph("4. Official Account Management Hierarchy (Government & Oil Sectors)")
    style_heading_1(h1)

    p = doc.add_paragraph(
        "In accordance with executive management directives, the account management structure for the Government and Oil Sectors "
        "is led by Marketing Manager Eiman Ashkanani, structured as follows:"
    )
    style_body(p)

    # Matrix Table
    table_matrix = doc.add_table(rows=7, cols=4)
    table_matrix.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_matrix.autofit = False

    matrix_headers = ["Sector", "Team Member", "Assigned Key Accounts", "Active Tenders Count"]
    hdr_row = table_matrix.rows[0]
    for j, h in enumerate(matrix_headers):
        cell = hdr_row.cells[j]
        set_cell_shading(cell, HEX_NAVY)
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.font.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)

    matrix_rows = [
        ("Gov & Oil Leadership", "Eiman Ashkanani\n(Marketing Manager)", "MOD (Defence), PAHW (Housing), MPW (Public Works), KM (Municipality), KPC, KIPIC, KUFPEC, KGOC, KAFCO, KPI, MOF, MOE, KU, PAI, KNG, KISR, KFAS, and all other public/oil accounts", "93 Tenders"),
        ("Government", "Jana Al-Obaid", "PIFSS (Social Security), PADA (Disability Affairs), Nazaha (Anti-Corruption), MOH (Ministry of Health)", "40 Tenders"),
        ("Government", "Raed Obeid", "MEW (Ministry of Electricity & Water), PACI (Civil Information)", "20 Tenders"),
        ("Government", "Khaled Alabdallah", "MOI (Ministry of Interior / الداخلية)", "9 Tenders"),
        ("Oil Sector", "Ahmed Habib", "KOC (Kuwait Oil Company), KOTC (Kuwait Oil Tanker Company)", "5 Tenders"),
        ("Oil Sector", "Abrar Al-Qallaf", "KNPC (Kuwait National Petroleum Company), PIC (Petrochemical Industries Company)", "64 Tenders")
    ]

    col_widths = [Inches(1.5), Inches(1.8), Inches(2.3), Inches(0.8)]
    for i, data in enumerate(matrix_rows):
        row = table_matrix.rows[i + 1]
        bg_col = HEX_LIGHT_GRAY if i % 2 == 0 else HEX_WHITE
        for j, val in enumerate(data):
            cell = row.cells[j]
            cell.width = col_widths[j]
            set_cell_shading(cell, bg_col)
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.size = Pt(8.5)
            if j == 1:
                run.font.bold = True
                run.font.color.rgb = COLOR_NAVY
            elif j == 3:
                run.font.bold = True
                run.font.color.rgb = COLOR_BLUE
            else:
                run.font.color.rgb = COLOR_TEXT

    p_after_t = doc.add_paragraph()
    p_after_t.paragraph_format.space_before = Pt(6)
    p_after_t.paragraph_format.space_after = Pt(8)

    # ================= 5. AUTOMATED DAILY REFRESH SCHEDULE =================
    h1 = doc.add_paragraph("5. Automated Daily Refresh Schedule")
    style_heading_1(h1)

    p = doc.add_paragraph(
        "To ensure zero lag between government gazette publication and sales pipeline action, the system executes an automated morning scan every day:"
    )
    style_body(p)

    p = doc.add_paragraph(
        "1. Schedule: 08:00 AM Kuwait Time (05:00 UTC) every morning.\n"
        "2. Cloud Automation Engine: GitHub Actions Scheduled Cron Workflow (.github/workflows/daily_scan.yml).\n"
        "3. Gazette Sync: Authenticates with Kuwait Al-Yawm, inspects latest gazette issues, downloads the tenders supplement PDF, and extracts text using pdfplumber.\n"
        "4. CAPT Sync: Scrapes the live opening and active tenders board from capt.gov.kw.\n"
        "5. Smart Incremental Merging: New tenders are seamlessly inserted with strategic tags; existing tenders receive updated closing dates; customized account owner assignments remain preserved.\n"
        "6. On-Demand Manual Refresh: Users can click 'Refresh Data' in the portal header or trigger a manual run on GitHub Actions at any time."
    )
    style_body(p)

    # ================= 6. GOVERNANCE & CONTACTS =================
    h1 = doc.add_paragraph("6. System Governance & Support")
    style_heading_1(h1)

    p = doc.add_paragraph(
        "For system support, feature requests, or new account mapping rules, please contact the KBM Technology Platform Team:\n"
        "• Production Platform: https://kbm-tenderscout.azurewebsites.net\n"
        "• Codebase & Issue Tracker: https://github.com/6439144/KBM_TenderScout\n"
        "• Platform Architecture: Python 3.12, FastAPI, Playwright Headless, Azure App Service (F1 Linux)"
    )
    style_body(p)

    add_callout(
        doc,
        [
            "Please bookmark the Azure URL in your daily browser favorites.",
            "Account managers are requested to review their filtered view each morning at 08:15 AM to capture upcoming bids well ahead of the pre-tender meeting deadlines."
        ],
        title="Best Practice Recommendation / توصيات الاستخدام الأمثل"
    )

    try:
        doc.save(str(OUTPUT_PATH))
        print(f"User Guide Word Document successfully saved at: {OUTPUT_PATH}")
    except PermissionError:
        alt_path = DOCS_DIR / "KBM_TenderScout_User_Guide_Updated.docx"
        doc.save(str(alt_path))
        print(f"Notice: Primary file was open in Word. Successfully saved updated version at: {alt_path}")

if __name__ == "__main__":
    build_guide()
