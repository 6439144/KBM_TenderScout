import openpyxl
import sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook("output/reports/KBM_Tenders_2026-10-05.xlsx")
print("Sheets:", wb.sheetnames)

ws_kbm = wb["KBM Opportunities"]
headers = [ws_kbm.cell(1, c).value for c in range(1, ws_kbm.max_column + 1)]
print("\nKBM Opportunities Headers:")
for i, h in enumerate(headers, 1):
    print(f"  Col {i}: {h}")

print("\nTop 5 rows in KBM Opportunities:")
for r in range(2, min(7, ws_kbm.max_row + 1)):
    row_vals = [ws_kbm.cell(r, c).value for c in range(1, ws_kbm.max_column + 1)]
    print(f"Row {r}: No={row_vals[0]} | Client={row_vals[1]} | Owner={row_vals[2]} | Title={row_vals[3][:40]}... | Scope={str(row_vals[4])[:50]}... | Fit={row_vals[6]}")

ws_all = wb["All Open"]
all_headers = [ws_all.cell(1, c).value for c in range(1, ws_all.max_column + 1)]
print("\nAll Open Headers:")
for i, h in enumerate(all_headers, 1):
    print(f"  Col {i}: {h}")
