import openpyxl
import sys
from pathlib import Path
import json

if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8")

file_path = r"C:\Users\kabed\OneDrive - KBM\Technology forecast - Technology forecast\2026-Q3-28th-Sep-Tech Pipeline Review-ver1.1.xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)
ws = wb["KBM Pipeline-2026"]

headers = [c.value for c in ws[1]]
sec_idx = headers.index("Sector")
am_idx = headers.index("AM")
acc_idx = headers.index("Account")

mapping = {}
for row in ws.iter_rows(min_row=2, values_only=True):
    acc = row[acc_idx]
    am = row[am_idx]
    sec = row[sec_idx]
    if acc and am:
        acc_clean = str(acc).strip()
        am_clean = str(am).strip()
        sec_clean = str(sec).strip() if sec else ""
        if acc_clean not in mapping:
            mapping[acc_clean] = {"am": am_clean, "sector": sec_clean, "count": 0}
        mapping[acc_clean]["count"] += 1

print(f"Total unique accounts: {len(mapping)}")
for acc, data in sorted(mapping.items()):
    print(f"Account: {acc:<35} | AM: {data['am']:<20} | Sector: {data['sector']}")

with open("data/account_managers.json", "w", encoding="utf-8") as f:
    json.dump(mapping, f, ensure_ascii=False, indent=2)
print("\nSaved mapping to data/account_managers.json")
